"""Helpers to build lecture decks on Christian's DTU template (lecture 13.1 master)."""
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
from lxml import etree

TPL = "/mnt/project-files/heating/lecture-13.1/13.1 Cost of floor heating and DHW - PEX refined.pptx"
RED, DARK, GREY, TINT, WHITE = RGBColor(0x99, 0, 0), RGBColor(0x26, 0x26, 0x26), RGBColor(0x6E, 0x6E, 0x6E), RGBColor(0xF2, 0xF2, 0xF2), RGBColor(255, 255, 255)
PINK = RGBColor(0xF6, 0xE6, 0xE6)
X0, X1, Y0, Y1 = Emu(1774726), Emu(1774726 + 9312374), Inches(1.87), Inches(6.85)
W = X1 - X0
FRONT_A, FRONT_B, ONE, TWO, TITLE_ONLY = 0, 1, 2, 3, 7


def _set(tf, text):
    p0 = tf.paragraphs[0]
    runs = p0.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def new_deck(footer, title):
    p = Presentation(TPL)
    lst = p.slides._sldIdLst
    for s in list(lst):
        p.part.drop_rel(s.rId); lst.remove(s)
    for sh in p.slide_masters[0].shapes:
        if sh.name == "date": _set(sh.text_frame, "41463 HPB")
        elif sh.name == "text" and "ZEB" in sh.text_frame.text: _set(sh.text_frame, footer)
    cp = p.core_properties
    cp.title, cp.author, cp.subject = title, "Christian Anker Hviid, DTU Construct", "41463 High Performance Buildings"
    p.core_properties.last_modified_by = "Christian Anker Hviid"
    return p


def layout(p, i):
    return p.slide_masters[0].slide_layouts[i]


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def title_slide(p, title, subtitle, note=""):
    s = p.slides.add_slide(layout(p, FRONT_A))
    s.placeholders[0].text = title; s.placeholders[1].text = subtitle
    if note: notes(s, note)
    return s


def section(p, title, kicker, note=""):
    s = p.slides.add_slide(layout(p, FRONT_B))
    s.placeholders[0].text = title; s.placeholders[1].text = kicker
    if note: notes(s, note)
    return s


def _bullet(par, level=0, char="•"):
    pPr = par._p.get_or_add_pPr()
    mar = 228600 + level * 228600
    pPr.set("marL", str(mar)); pPr.set("indent", str(-228600))
    for t in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(t)): pPr.remove(e)
    bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", char if level == 0 else "–")


def text(s, x, y, w, h, items, size=16, color=DARK, bullets=True, bold_first=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6, name=None):
    """items: list of str | (str, level) | list of (text, {bold, color, italic}) runs."""
    tb = s.shapes.add_textbox(x, y, w, h)
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0; tf.margin_top = tf.margin_bottom = Emu(0)
    for i, it in enumerate(items):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align; par.space_after = Pt(space)
        level = 0
        if isinstance(it, tuple) and isinstance(it[1], int): it, level = it
        runs = it if isinstance(it, list) else [(it, {})]
        for t, o in runs:
            r = par.add_run(); r.text = t
            f = r.font; f.size = Pt(o.get("size", size - 2 * level)); f.color.rgb = o.get("color", color)
            f.bold = o.get("bold", bold_first and i == 0); f.italic = o.get("italic", False)
            if o.get("font"): f.name = o["font"]
        if bullets: _bullet(par, level)
    return tb


def title_only(p, title, note=""):
    s = p.slides.add_slide(layout(p, TITLE_ONLY))
    s.shapes.title.text = title
    if note: notes(s, note)
    return s


def bullets_slide(p, title, items, note=""):
    s = p.slides.add_slide(layout(p, ONE))
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    for i, it in enumerate(items):
        level = 0
        if isinstance(it, tuple): it, level = it
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = it; par.level = level
    if note: notes(s, note)
    return s


def picture(s, path, x, y, w, h, align="center"):
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih); pw, ph = int(iw * sc), int(ih * sc)
    px = x + (w - pw) // 2 if align == "center" else (x if align == "left" else x + w - pw)
    pic = s.shapes.add_picture(path, px, y + (h - ph) // 2, pw, ph)
    return pic


def caption(s, x, y, w, t):
    return text(s, x, y, w, Inches(0.4), [t], size=11, color=GREY, bullets=False)


def box(s, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False, name=None):
    b = s.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = 0.06
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(1.5)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            b.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    b.shadow.inherit = False
    if name: b.name = name
    b.text_frame.text = ""
    return b


def card(s, x, y, w, h, head, body, fill=TINT, head_color=RED, size=14, name=None):
    box(s, x, y, w, h, fill=fill, name=name)
    pad = Inches(0.18)
    text(s, x + pad, y + pad, w - 2 * pad, Inches(0.45), [head], size=size + 3, color=head_color, bullets=False, bold_first=True)
    if body:
        text(s, x + pad, y + pad + Inches(0.5), w - 2 * pad, h - Inches(0.7) - pad, body, size=size)


def equation(s, x, y, w, h, eq, label=None, size=20):
    box(s, x, y, w, h, fill=PINK)
    items = []
    if label: items.append([(label, {"size": 12, "color": GREY})])
    items.append([(eq, {"size": size, "font": "Cambria", "italic": False})])
    text(s, x + Inches(0.2), y + Inches(0.1), w - Inches(0.4), h - Inches(0.2), items, bullets=False, anchor=MSO_ANCHOR.MIDDLE, space=2)


def todo(s, x, y, w, h, lines):
    box(s, x, y, w, h, fill=None, line=GREY, dash=True, shape=MSO_SHAPE.RECTANGLE, name="To develop")
    text(s, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), h - Inches(0.4),
         [[("TO DEVELOP", {"bold": True, "color": RED, "size": 14})]] + lines, size=16, color=GREY, bullets=False, space=8)


def stat(s, x, y, w, big, small, color=RED, size=40):
    text(s, x, y, w, Inches(0.8), [[(big, {"size": size, "bold": True, "color": color})]], bullets=False)
    text(s, x, y + Inches(0.85), w, Inches(0.7), [small], size=13, color=GREY, bullets=False)


PARTS = [("1", "Window size", "AutoMOO"), ("2", "Venting hatch", "jet theory, draught"),
         ("3", "Ceiling fans", "peak periods"), ("4", "Shutters", "heat wave")]


def strategy_strip(s, active, y=Inches(2.2)):
    n = len(PARTS); gap = Inches(0.35); w = (W - gap * (n - 1)) // n; h = Inches(1.5)
    for i, (num, head, sub) in enumerate(PARTS):
        x = X0 + i * (w + gap); on = i == active
        b = box(s, x, y, w, h, fill=RED if on else TINT, name=f"Part {num}")
        c = WHITE if on else DARK
        text(s, x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), h - Inches(0.3),
             [[(f"Part {num}", {"size": 13, "color": WHITE if on else GREY})], [(head, {"size": 19, "bold": True, "color": c})],
              [(sub, {"size": 13, "color": c})]], bullets=False, space=2)
        if i < n - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + w + Inches(0.05), y + h // 2 - Inches(0.12), gap - Inches(0.1), Inches(0.24))
            a.fill.solid(); a.fill.fore_color.rgb = GREY; a.line.fill.background(); a.shadow.inherit = False


def flow(s, steps, y, h=Inches(1.1), x=X0, w=W, size=15, active=None):
    n = len(steps); gap = Inches(0.4); bw = (w - gap * (n - 1)) // n
    for i, st in enumerate(steps):
        bx = x + i * (bw + gap); on = active == i
        box(s, bx, y, bw, h, fill=RED if on else TINT)
        lines = st if isinstance(st, list) else [st]
        text(s, bx + Inches(0.12), y + Inches(0.08), bw - Inches(0.24), h - Inches(0.16),
             [[(l, {"bold": j == 0, "color": WHITE if on else DARK, "size": size if j == 0 else size - 2})] for j, l in enumerate(lines)],
             bullets=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=2)
        if i < n - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, bx + bw + Inches(0.06), y + h // 2 - Inches(0.12), gap - Inches(0.12), Inches(0.24))
            a.fill.solid(); a.fill.fore_color.rgb = GREY; a.line.fill.background(); a.shadow.inherit = False


def where_we_are(p, active, line, note=""):
    s = title_only(p, "Where we are in the cooling assignment", note)
    strategy_strip(s, active)
    text(s, X0, Inches(4.2), W, Inches(1.5), line, size=18, bullets=True)
    return s


def fig_slide(p, title, img, bullets=None, cap=None, note="", split=0.42):
    s = title_only(p, title, note)
    if bullets:
        tw = int(W * split)
        text(s, X0, Y0, tw - Inches(0.3), Y1 - Y0, bullets, size=16)
        picture(s, img, X0 + tw, Y0, W - tw, Y1 - Y0 - (Inches(0.45) if cap else 0))
        if cap: caption(s, X0 + tw, Y1 - Inches(0.4), W - tw, cap)
    else:
        picture(s, img, X0, Y0, W, Y1 - Y0 - (Inches(0.45) if cap else 0))
        if cap: caption(s, X0, Y1 - Inches(0.4), W, cap)
    return s


def table(s, x, y, w, rows, col_w=None, size=13, row_h=Inches(0.42)):
    t = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows)).table
    if col_w:
        for i, cw in enumerate(col_w): t.columns[i].width = int(w * cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c); cell.text = val
            cell.fill.solid(); cell.fill.fore_color.rgb = RED if r == 0 else (TINT if r % 2 == 0 else WHITE)
            for par in cell.text_frame.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(size); run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else DARK
            cell.margin_left = cell.margin_right = Inches(0.08)
    return t
